# ruff: noqa: E501
"""Parse CSV/XLSX/PDF endpoint - AI-powered batch transaction extraction."""

from __future__ import annotations

import base64
import csv
import json
from datetime import date, datetime
from io import BytesIO, StringIO
from pathlib import Path
from typing import Any, cast

from fastapi import APIRouter, File, UploadFile
from fastapi.responses import JSONResponse
from langchain_core.messages import HumanMessage
from loguru import logger
from openpyxl import load_workbook
from pypdf import PdfReader

from ..config import ALLOWED_CATEGORIES, OPENAI_API_KEY
from ..schemas import ParseBatchResponse
from ..services.ai import llm

router = APIRouter()

MAX_BATCH_FILE_SIZE = 5 * 1024 * 1024
MAX_BATCH_TRANSACTIONS = 100
CSV_EXTENSIONS = {".csv"}
EXCEL_EXTENSIONS = {".xlsx"}
PDF_EXTENSIONS = {".pdf"}
IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".webp", ".gif"}
ALLOWED_CONTENT_TYPES = {
    "text/csv",
    "application/csv",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "application/pdf",
    "image/jpeg",
    "image/png",
    "image/webp",
    "image/gif",
}
HEADER_HINTS = {
    "amount",
    "category",
    "date",
    "description",
    "expense",
    "income",
    "note",
    "price",
    "transaction",
    "type",
    "value",
    "chi",
    "danh mục",
    "dien giai",
    "diễn giải",
    "ghi chú",
    "loại",
    "mô tả",
    "ngày",
    "số tiền",
    "thu",
}

BATCH_PARSING_PROMPT = """
<identity>
You are a financial transaction parser. Convert CSV, Excel, or PDF bank-statement content into structured JSON transactions.
</identity>

<context>
Today's date: {today}
Input language: {lang_hint}
Valid categories: {categories_list}
Maximum transactions to return: {max_transactions}
</context>

<rows_json>
{rows_json}
</rows_json>

<rules>
1. Return JSON with a top-level "transactions" array only.
2. Each item must contain exactly:
   - "type": "expense" or "income"
   - "amount": positive number
   - "category": one of {categories_list}
   - "description": concise string
   - "transactionDate": YYYY-MM-DD string
3. Accept any column naming style, bank statement layout, and PDF text ordering. Infer field meaning from headers and row content.
4. Skip rows that are totals, balances, empty, duplicated headers, notes, or cannot be confidently converted to a transaction.
5. Use "expense" by default. Use "income" only when row content clearly indicates money received: lương, nhận, thưởng, thu nhập, salary, income, received, deposit, credit.
6. Preserve numeric amounts accurately. Interpret Vietnamese shorthand when present: k = 1,000; tr/triệu/củ/chai = 1,000,000; lít/lốp/sọi = 100,000; xị/xịch = 10,000.
7. Normalize dates to YYYY-MM-DD. If a row has no date, use today's date ({today}).
8. Category must be exactly one valid category. Use "Khác" if uncertain.
9. Return at most {max_transactions} transactions.
</rules>
"""

BATCH_IMAGE_PROMPT = """
<identity>
You are a financial transaction parser. Convert this image into structured JSON transactions.
</identity>

<context>
Today's date: {today}
Input language: {lang_hint}
Valid categories: {categories_list}
Maximum transactions to return: {max_transactions}
</context>

<rules>
1. Return JSON with a top-level "transactions" array only.
2. Extract every visible transaction row, receipt total, or statement line that can be confidently parsed.
3. Each item must contain exactly:
   - "type": "expense" or "income"
   - "amount": positive number
   - "category": one of {categories_list}
   - "description": concise string
   - "transactionDate": YYYY-MM-DD string
4. Skip totals, balances, headers, duplicated rows, notes, and unreadable lines.
5. Use "expense" by default. Use "income" only when row content clearly indicates money received: lương, nhận, thưởng, thu nhập, salary, income, received, deposit, credit.
6. Preserve numeric amounts accurately. Interpret Vietnamese number formatting and shorthand: k = 1,000; tr/triệu/củ/chai = 1,000,000; lít/lốp/sọi = 100,000; xị/xịch = 10,000.
7. Normalize dates to YYYY-MM-DD. If a transaction has no date, use today's date ({today}).
8. Category must be exactly one valid category. Use "Khác" if uncertain.
9. Return at most {max_transactions} transactions.
</rules>

Respond in JSON format.
"""


def _decode_csv(raw: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "cp1258", "latin-1"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def _normalize_cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def _rows_from_csv(raw: bytes) -> list[dict[str, str]]:
    text = _decode_csv(raw)
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample)
    except csv.Error:
        dialect = csv.excel

    reader = csv.reader(StringIO(text), dialect)
    rows = [[cell.strip() for cell in row] for row in reader if any(cell.strip() for cell in row)]
    return _rows_to_dicts(rows)


def _rows_from_xlsx(raw: bytes) -> list[dict[str, str]]:
    workbook = load_workbook(filename=BytesIO(raw), read_only=True, data_only=True)
    sheet = workbook.active
    rows: list[list[str]] = []
    try:
        for row in sheet.iter_rows(values_only=True):
            values = [_normalize_cell(cell) for cell in row]
            if any(values):
                rows.append(values)
            if len(rows) > MAX_BATCH_TRANSACTIONS + 1:
                break
    finally:
        workbook.close()
    return _rows_to_dicts(rows)


def _rows_from_pdf(raw: bytes) -> list[dict[str, str]]:
    reader = PdfReader(BytesIO(raw))
    rows: list[dict[str, str]] = []

    for page_index, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        lines = [line.strip() for line in text.splitlines() if line.strip()]
        for line_index, line in enumerate(lines, start=1):
            rows.append(
                {
                    "page": str(page_index),
                    "line": str(line_index),
                    "text": line,
                }
            )
            if len(rows) >= MAX_BATCH_TRANSACTIONS * 3:
                return rows

    return rows


def _is_image_upload(extension: str, content_type: str) -> bool:
    return extension in IMAGE_EXTENSIONS or content_type.startswith("image/")


def _rows_to_dicts(rows: list[list[str]]) -> list[dict[str, str]]:
    if not rows:
        return []

    header_index = 0
    has_header = False
    for index, row in enumerate(rows[:10]):
        non_empty = [cell for cell in row if cell]
        if len(non_empty) >= 2:
            header_index = index
            normalized = " ".join(non_empty).lower()
            has_header = any(hint in normalized for hint in HEADER_HINTS)
            break

    max_columns = max(len(row) for row in rows)
    headers = [f"column_{index + 1}" for index in range(max_columns)]
    start_index = 0

    if has_header:
        headers = [
            value if value else f"column_{index + 1}"
            for index, value in enumerate(rows[header_index])
        ]
        start_index = header_index + 1

    output: list[dict[str, str]] = []

    for row in rows[start_index:]:
        item = {
            headers[index] if index < len(headers) else f"column_{index + 1}": value
            for index, value in enumerate(row)
            if value
        }
        if item:
            output.append(item)
        if len(output) >= MAX_BATCH_TRANSACTIONS:
            break

    return output


@router.post("/api/parse-batch", response_model=ParseBatchResponse)
async def parse_batch(
    file: UploadFile = File(...),  # noqa: B008
    language: str = "vi",
) -> ParseBatchResponse | JSONResponse:
    """Parse CSV/XLSX/PDF content into structured transactions."""

    if not OPENAI_API_KEY:
        return JSONResponse(
            status_code=503,
            content={"error": "AI service not configured", "redirect": "/transactions"},
        )

    filename = file.filename or ""
    extension = Path(filename).suffix.lower()
    content_type = file.content_type or ""

    if (
        extension not in CSV_EXTENSIONS | EXCEL_EXTENSIONS | PDF_EXTENSIONS | IMAGE_EXTENSIONS
        and content_type not in ALLOWED_CONTENT_TYPES
    ):
        return JSONResponse(
            status_code=400,
            content={"error": "Only CSV, XLSX, PDF, and image files are supported"},
        )

    raw = await file.read(MAX_BATCH_FILE_SIZE + 1)
    if len(raw) > MAX_BATCH_FILE_SIZE:
        return JSONResponse(
            status_code=400, content={"error": "File too large. Maximum size is 5 MB"}
        )

    today = date.today().isoformat()
    lang_hint = "Vietnamese" if language == "vi" else "English"
    categories_list = ", ".join(ALLOWED_CATEGORIES)

    try:
        structured_llm = llm.with_structured_output(ParseBatchResponse)

        if _is_image_upload(extension, content_type):
            base64_image = base64.b64encode(raw).decode("utf-8")
            prompt = BATCH_IMAGE_PROMPT.format(
                today=today,
                lang_hint=lang_hint,
                categories_list=categories_list,
                max_transactions=MAX_BATCH_TRANSACTIONS,
            )
            message = HumanMessage(
                content=[
                    {"type": "text", "text": prompt},
                    {
                        "type": "image_url",
                        "image_url": {"url": f"data:{content_type};base64,{base64_image}"},
                    },
                ]
            )
            result = structured_llm.invoke([message])
            return cast(ParseBatchResponse, result)

        try:
            if extension in EXCEL_EXTENSIONS or content_type.endswith("spreadsheetml.sheet"):
                rows = _rows_from_xlsx(raw)
            elif extension in PDF_EXTENSIONS or content_type == "application/pdf":
                rows = _rows_from_pdf(raw)
            else:
                rows = _rows_from_csv(raw)
        except Exception as exc:
            logger.error("Batch file read error: %s", exc)
            return JSONResponse(status_code=400, content={"error": "Failed to read file"})

        if not rows:
            return JSONResponse(status_code=400, content={"error": "No transaction rows found"})

        prompt = BATCH_PARSING_PROMPT.format(
            today=today,
            lang_hint=lang_hint,
            categories_list=categories_list,
            rows_json=json.dumps(rows, ensure_ascii=False),
            max_transactions=MAX_BATCH_TRANSACTIONS,
        )
        result = structured_llm.invoke(prompt)
        return cast(ParseBatchResponse, result)

    except Exception as exc:
        logger.error("AI batch parse error: %s", exc)
        return JSONResponse(
            status_code=500,
            content={"error": "Failed to parse transactions", "redirect": "/transactions"},
        )
